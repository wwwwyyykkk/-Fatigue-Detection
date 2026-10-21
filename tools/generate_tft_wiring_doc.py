from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt


OUTPUT = "output/doc/ESP32-S3_TFT_接线文档.docx"


def set_font(run, size=10.5, bold=False):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(document, text, size=10.5, bold=False):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold)
    return paragraph


def fill_cell(cell, text, bold=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.runs.clear()
    run = paragraph.add_run(text)
    set_font(run, size=9, bold=bold)


def add_table(document, title, headers, rows):
    add_paragraph(document, title, size=12, bold=True)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        fill_cell(table.rows[0].cells[index], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            fill_cell(cells[index], value)
    document.add_paragraph()
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ESP32-S3 疲劳检测项目 TFT 版接线文档")
    set_font(run, size=18, bold=True)

    add_paragraph(doc, "版本说明：本表以你之前的 OLED 接线为基础，将 SSD1306 OLED 替换为 SPI TFT 显示屏。摄像头引脚保持不变，蜂鸣器、MQ-3、EG800K、按键和电源接线保留。", size=10.5)
    add_paragraph(doc, "重要原则：摄像头已占用 GPIO2、4、5、6、7、8、9、10、11、12、13、14、15、16、17、18，其他模块不要再使用这些引脚。", size=10.5, bold=True)

    all_rows = [
        ("OV5640 摄像头", "PWDN", "GPIO15", "输出", "摄像头电源管理"),
        ("OV5640 摄像头", "RESET", "GPIO5", "输出", "摄像头复位"),
        ("OV5640 摄像头", "XCLK", "GPIO14", "输出", "摄像头时钟"),
        ("OV5640 摄像头", "SIOD / SDA", "GPIO2", "双向", "SCCB / I2C 数据"),
        ("OV5640 摄像头", "SIOC / SCL", "GPIO4", "输出", "SCCB / I2C 时钟"),
        ("OV5640 摄像头", "Y7", "GPIO16", "输入", "图像数据 D7"),
        ("OV5640 摄像头", "Y6", "GPIO17", "输入", "图像数据 D6"),
        ("OV5640 摄像头", "Y5", "GPIO18", "输入", "图像数据 D5"),
        ("OV5640 摄像头", "Y4", "GPIO12", "输入", "图像数据 D4"),
        ("OV5640 摄像头", "Y3", "GPIO10", "输入", "图像数据 D3"),
        ("OV5640 摄像头", "Y2", "GPIO8", "输入", "图像数据 D2"),
        ("OV5640 摄像头", "Y1", "GPIO9", "输入", "图像数据 D1"),
        ("OV5640 摄像头", "Y0", "GPIO11", "输入", "图像数据 D0"),
        ("OV5640 摄像头", "VSYNC", "GPIO6", "输入", "场同步"),
        ("OV5640 摄像头", "HREF", "GPIO7", "输入", "行同步"),
        ("OV5640 摄像头", "PCLK", "GPIO13", "输入", "像素时钟"),
        ("SPI TFT 显示屏", "SCL / SCK / CLK", "GPIO40", "输出", "SPI 时钟"),
        ("SPI TFT 显示屏", "SDA / MOSI / DIN", "GPIO41", "输出", "SPI 数据输出"),
        ("SPI TFT 显示屏", "CS", "GPIO42", "输出", "片选"),
        ("SPI TFT 显示屏", "DC / A0", "GPIO39", "输出", "数据/命令选择"),
        ("SPI TFT 显示屏", "RST / RES", "3.3V 或 GPIO47", "输出/固定", "建议接 3.3V 并在代码中设为不使用复位脚；若必须软件复位，用 GPIO47"),
        ("SPI TFT 显示屏", "BL / LED / BLK", "GPIO21 或 3.3V", "输出/电源", "需要亮度控制接 GPIO21；不需要控制直接接 3.3V"),
        ("SPI TFT 显示屏", "MISO / SDO", "不接", "-", "只显示不读屏，可不接"),
        ("无源蜂鸣器", "PWM", "GPIO38", "输出", "保留原接线，PWM 控制报警"),
        ("MQ-3 酒精传感器", "AO", "GPIO1", "输入", "ADC1_0，模拟量输入"),
        ("EG800K 4G+GPS", "TX -> ESP RX", "GPIO44", "输入", "UART0 RX / U0RXD"),
        ("EG800K 4G+GPS", "RX <- ESP TX", "GPIO43", "输出", "UART0 TX / U0TXD"),
        ("按键 1", "信号", "GPIO3", "输入", "行程开始/结束，建议上拉"),
        ("按键 2", "信号", "GPIO46", "输入", "疲劳确认；注意 GPIO46 为启动相关引脚，启动后可正常使用"),
        ("系统电源", "GND", "GND", "电源", "所有模块必须共地"),
        ("系统电源", "3.3V", "3V3", "电源", "TFT 逻辑、电平和可选背光；MQ-3 AO 不能超过 3.3V"),
        ("系统电源", "5V", "5V0", "电源", "蜂鸣器、MQ-3 模块可用，视模块规格决定"),
        ("EG800K 电源", "外部稳压 5-12V", "独立电源", "电源", "EG800K 电流大，建议独立供电，不从板载 5V 直接取电"),
    ]

    add_table(doc, "一、完整引脚分配总表", ["模块", "接口/信号", "ESP32-S3 接口", "方向/类型", "说明"], all_rows)

    tft_rows = [
        ("VCC", "3.3V", "不要用 5V 信号电平；若模块标注支持 5V 供电，也建议先用 3.3V"),
        ("GND", "GND", "必须与 ESP32-S3 共地"),
        ("SCL / SCK / CLK", "GPIO40", "SPI 时钟"),
        ("SDA / MOSI / DIN", "GPIO41", "SPI 数据"),
        ("CS", "GPIO42", "片选"),
        ("DC / A0", "GPIO39", "数据/命令选择"),
        ("RST / RES", "3.3V 或 GPIO47", "优先接 3.3V；代码里将 RST 设为 -1。需要软件复位时才接 GPIO47"),
        ("BL / LED / BLK", "GPIO21 或 3.3V", "直接接 3.3V 最简单；接 GPIO21 可 PWM 调亮度"),
        ("MISO / SDO", "不接", "显示场景不需要读屏"),
    ]
    add_table(doc, "二、TFT 显示屏单独接线表", ["TFT 引脚", "连接到", "备注"], tft_rows)

    old_new_rows = [
        ("SSD1306 OLED SDA", "GPIO41", "取消", "改为 TFT MOSI / SDA / DIN"),
        ("SSD1306 OLED SCL", "GPIO42", "取消", "改为 TFT CS"),
        ("新增 TFT SCLK", "GPIO40", "新增", "SPI 时钟"),
        ("新增 TFT DC", "GPIO39", "新增", "命令/数据选择"),
        ("新增 TFT RST", "3.3V 或 GPIO47", "新增", "不占用蜂鸣器 GPIO38"),
        ("新增 TFT BL", "GPIO21 或 3.3V", "新增", "背光供电或亮度控制"),
    ]
    add_table(doc, "三、OLED 更换为 TFT 后的变化", ["原接口/新增接口", "原 GPIO/新连接", "状态", "说明"], old_new_rows)

    code_rows = [
        ("TFT_SCLK", "40"),
        ("TFT_MOSI", "41"),
        ("TFT_CS", "42"),
        ("TFT_DC", "39"),
        ("TFT_RST", "-1 或 47"),
        ("TFT_BL", "21"),
    ]
    add_table(doc, "四、建议代码宏定义", ["宏名", "值"], code_rows)

    add_paragraph(doc, "建议宏定义：", size=11, bold=True)
    code = (
        "#define TFT_SCLK 40\n"
        "#define TFT_MOSI 41\n"
        "#define TFT_CS   42\n"
        "#define TFT_DC   39\n"
        "#define TFT_RST  -1   // RST 接 3.3V 时使用；若接 GPIO47 则改为 47\n"
        "#define TFT_BL   21"
    )
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(code)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Consolas")
    run.font.size = Pt(10)

    add_paragraph(doc, "检查项：", size=11, bold=True)
    checks = [
        "所有模块 GND 必须共地。",
        "TFT 信号线使用 3.3V 逻辑。",
        "不要把 TFT 接到摄像头已占用的 GPIO2、4、5、6、7、8、9、10、11、12、13、14、15、16、17、18。",
        "GPIO38 已保留给无源蜂鸣器，不再分配给 TFT RST。",
        "EG800K 建议独立供电，避免启动瞬间电流导致 ESP32-S3 复位。",
        "GPIO46 是启动相关引脚，按键电路不要在上电启动时强制拉到错误电平。",
    ]
    for item in checks:
        add_paragraph(doc, f"- {item}", size=10.5)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
